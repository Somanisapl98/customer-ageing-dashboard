from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime,date
import json
ROOT=Path(__file__).resolve().parents[1]
def read(path,header):
 wb=load_workbook(path,read_only=True,data_only=True); ws=wb[wb.sheetnames[0]]; heads=[str(x.value or '').strip() for x in ws[header]]; out=[]
 for vals in ws.iter_rows(min_row=header+1,values_only=True):
  d={}
  for i,h in enumerate(heads):
   if not h: continue
   v=vals[i] if i<len(vals) else None
   if isinstance(v,(datetime,date)):v=v.strftime('%Y-%m-%d')
   d[h]=v
  if any(v not in (None,'') for v in vals):out.append(d)
 return out
age=[r for r in read(ROOT/'input'/'Customer Aging Analyses Bill To Bill.xlsx',4) if r.get('Bill No') and r.get('party_code')]
master=[r for r in read(ROOT/'input'/'Customer Master.xlsx',4) if r.get('Code') and r.get('Customer')]
wa=read(ROOT/'input'/'Customer WhatsApp Master FINAL.xlsx',1)
(ROOT/'data'/'dashboard-data.json').write_text(json.dumps({'generatedAt':datetime.now().isoformat(timespec='seconds'),'ageing':age,'master':master,'whatsapp':wa},ensure_ascii=False),encoding='utf-8')
print('Dashboard data updated.')
